"""Provider -> model -> status -> report/provenance integration without network."""
import copy
import json
from datetime import datetime,timezone
from pathlib import Path
import pandas as pd
import pytest
from openpyxl import load_workbook
from screener import __main__ as cli,market
from test_metrics import operating,FX

@pytest.mark.parametrize('bank_only',[False,True])
def test_cli_writes_reports_and_provenance_even_when_every_name_abstains(tmp_path,monkeypatch,bank_only):
    config=(Path(__file__).parents[1]/'config.toml').read_text(encoding='utf-8').replace('pause_seconds = 0.4','pause_seconds = 0.0')
    (tmp_path/'config.toml').write_text(config,encoding='utf-8')
    (tmp_path/'input').mkdir()
    (tmp_path/'input/watchlist.csv').write_text('symbol,name,yahoo_symbol\nTEST,Test,TEST\n',encoding='utf-8')
    d=operating();d['info']['country']='United States'
    if bank_only:d['info']['industry']='Banks - Regional'
    stamp=datetime.now(timezone.utc)
    d.update(fetched_at=stamp,source_snapshot='controlled-test-record',source_sha256='controlled')
    d['info']['regularMarketTime']=int(stamp.timestamp())
    quote=market.Observation(20,stamp.isoformat(),'controlled quote',d['info'])
    monkeypatch.setattr(cli,'ROOT',tmp_path)
    monkeypatch.setattr(cli.fetch,'get',lambda *a,**k:(copy.deepcopy(d),True))
    monkeypatch.setattr(cli.fetch,'latest_prices',lambda symbols:{'TEST':quote})
    monkeypatch.setattr(cli.fetch,'FX',FX)
    assert cli.main([])==0
    csv=next((tmp_path/'data/history').glob('screen_*.csv'))
    frame=pd.read_csv(csv)
    assert len(frame)==1 and frame.iloc[0]['class']!='error'
    if bank_only:assert frame.iloc[0]['status']=='not rankable'
    receipt=json.loads(next((tmp_path/'data/history/provenance').glob('*.json')).read_text())
    assert receipt['implementation_sha256']['valuation.py']
    assert receipt['configuration']['valuation']['years']==10
    assert frame.iloc[0]['valuation_date']==receipt['date']
    book=load_workbook(next((tmp_path/'reports').glob('*.xlsx')),read_only=True)
    assert {'Candidates','Operating','Financials','Not rankable','Method'}<=set(book.sheetnames)
    book.close()
